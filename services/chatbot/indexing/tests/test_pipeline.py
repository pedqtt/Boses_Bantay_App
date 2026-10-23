import json
import sys
import unittest
import tempfile
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'server'))
from config import Settings,ROOT
from engine import Engine,select_context,numbers
from errors import PipelineFailure
from ingestion import chunk_document,source_units


def row(i=1,doc=1,text='Applications cost 25 pesos.',score=.95):
    return dict(chunk_id=i,document_id=doc,revision=1,title=f'Source {doc}',section_title='Requirements',content=text,
                content_hash=str(i),token_count=30,provenance={},similarity=.8,relevance=score,metadata={})

class DB:
    def __init__(self,rows=None,indexed=1,approved=1,valid=True):
        self.rows=rows if rows is not None else [row()];self.indexed=indexed;self.approved=approved;self.valid=valid
    def catalogue(self):return {'approved':self.approved,'indexed':self.indexed,'titles':['Source 1']}
    def search(self,*args):return self.rows
    def validate(self,*args):return self.valid

class Models:
    def __init__(self,draft=None,fail=None,entail=.99,relevance=.95):
        self.draft=draft or {'status':'answered','claims':[{'text':'Applications cost 25 pesos.','evidence':[{'source_id':1,'quote':'Applications cost 25 pesos.'}]}]}
        self.fail=fail;self.entail=entail;self.relevance=relevance
    def embed(self,texts):
        if self.fail=='embed':raise PipelineFailure('embed')
        return [[0.1]*1024 for _ in texts]
    def call(self,op,payload):
        if self.fail==op:raise PipelineFailure(op)
        if op=='generate':
            if 'missing_requested_details' in payload.get('schema',{}).get('properties',{}):
                return {'text':json.dumps({'missing_requested_details':self.draft.get('missing_aspects',[])})}
            if 'search_queries' in payload.get('schema',{}).get('properties',{}):
                query=json.loads(payload['prompt'])['question']
                return {'text':json.dumps({'query':query,'aspects':[query],'search_queries':[query],'subject_quote':query})}
            return {'text':json.dumps(self.draft)}
        if op=='rerank':return {'scores':[self.relevance]*len(payload['pairs'])}
        if op=='verify':return {'scores':[{'entailment':self.entail,'contradiction':1-self.entail} for _ in payload['pairs']]}

class PipelineTests(unittest.TestCase):
    def engine(self,db=None,models=None):return Engine(Settings('',''),db or DB(),models or Models())
    def test_answer_cites_checked_source(self):
        result=self.engine().answer('What is the application fee?')
        self.assertEqual(result['status'],'answered');self.assertIn('[1]',result['answer']);self.assertEqual(len(result['sources']),1)
    def test_wrong_number_rejected(self):
        m=Models();m.draft['claims'][0]['text']='Applications cost 250 pesos.'
        self.assertEqual(self.engine(models=m).answer('Fee?')['status'],'low_confidence')
    def test_invented_quote_rejected(self):
        m=Models();m.draft['claims'][0]['evidence'][0]['quote']='Applications are free.'
        self.assertEqual(self.engine(models=m).answer('Fee?')['status'],'low_confidence')
    def test_unsupported_negation_rejected(self):
        self.assertEqual(self.engine(models=Models(entail=.1)).answer('Fee?')['status'],'low_confidence')
    def test_no_information_is_question_specific(self):
        r=self.engine(db=DB(rows=[])).answer('Passport delivery time?')
        self.assertEqual(r['status'],'no_information');self.assertIn('Passport delivery time',r['answer']);self.assertFalse(r['retryable'])
    def test_low_confidence_distinct_from_no_information(self):
        r=self.engine(models=Models(relevance=.05)).answer('Fee?')
        self.assertEqual(r['status'],'low_confidence')
    def test_model_failures_are_technical(self):
        for stage in ['embed','rerank','generate','verify']:
            r=self.engine(models=Models(fail=stage)).answer('Fee?')
            self.assertEqual((r['status'],r['failureStage']),('service_failure',stage))
            self.assertNotIn('could not find',r['answer']);self.assertTrue(r['retryable'])
    def test_backlog_not_missing_content(self):
        r=self.engine(db=DB(indexed=0)).answer('Fee?')
        self.assertEqual(r['failureStage'],'indexing')
    def test_withdrawn_source_not_served(self):
        r=self.engine(db=DB(valid=False)).answer('Fee?')
        self.assertEqual(r['failureStage'],'source_validation');self.assertEqual(r['sources'],[])
    def test_partial_identifies_missing_aspect(self):
        m=Models();m.draft.update(status='partial',missing_aspects=['processing time'])
        r=self.engine(models=m).answer('Fee and processing time?')
        self.assertEqual(r['status'],'partial');self.assertIn('processing time',r['answer'])
    def test_coverage_review_cannot_erase_an_identified_gap(self):
        class OverconfidentReview(Models):
            def call(self,op,payload):
                if op=='generate' and 'missing_requested_details' in payload.get('schema',{}).get('properties',{}):
                    return {'text':json.dumps({'missing_requested_details':[]})}
                return super().call(op,payload)
        m=OverconfidentReview();m.draft.update(status='partial',missing_aspects=['processing time'])
        r=self.engine(models=m).answer('Fee and processing time?')
        self.assertEqual(r['status'],'partial');self.assertIn('processing time',r['answer'])
    def test_ambiguity_uses_source_options(self):
        m=Models({'status':'ambiguous','clarification':'Which application do you mean?','options':['Invented service']})
        r=self.engine(models=m).answer('How do I apply?')
        self.assertEqual(r['status'],'ambiguous');self.assertIn('Source 1',r['answer']);self.assertNotIn('Invented service',r['answer'])
    def test_invalid_conflict_downgraded(self):
        m=Models({'status':'conflict','disagreements':[]})
        self.assertEqual(self.engine(models=m).answer('Fee?')['status'],'low_confidence')
    def test_context_deduplicates_and_preserves_multisource(self):
        rows=[row(),row(2,text='Applications cost 25 pesos.'),row(3,doc=2,text='Processing takes three days.')]
        context=select_context(rows,Settings('',''))
        self.assertEqual(len(context),2);self.assertEqual({r['document_id'] for r in context},{1,2})
    def test_capacity_is_retryable_failure(self):
        e=self.engine()
        for _ in range(4):self.assertTrue(e.admission.acquire(False))
        self.assertEqual(e.answer('Fee?')['failureStage'],'capacity')
    def test_normalized_numbers(self):self.assertEqual(numbers('1,250.50 and 2%'),{'1250.50','2%'})
    def test_partial_index_still_answers_with_coverage_notice(self):
        r=self.engine(db=DB(indexed=1,approved=2)).answer('Application fee?')
        self.assertEqual(r['status'],'answered');self.assertTrue(r['coverageIncomplete'])
    def test_database_outage_is_not_content_absence(self):
        class Offline(DB):
            def search(self,*args):raise PipelineFailure('retrieval')
        r=self.engine(db=Offline()).answer('Application fee?')
        self.assertEqual(r['failureStage'],'retrieval');self.assertNotIn('could not find',r['answer'])
    def test_nonfinite_relevance_fails_closed(self):
        r=self.engine(models=Models(relevance=float('nan'))).answer('Application fee?')
        self.assertEqual(r['status'],'service_failure')
    def test_unknown_subject_with_multiple_sources_requests_clarification(self):
        class Unknown(Models):
            def call(self,op,payload):
                r=super().call(op,payload)
                if op=='generate' and 'search_queries' in payload.get('schema',{}).get('properties',{}):
                    plan=json.loads(r['text']);plan['subject_quote']='';r['text']=json.dumps(plan)
                return r
        r=self.engine(db=DB(rows=[row(),row(2,doc=2,text='Other service requirements.')]),models=Unknown()).answer('How do I apply?')
        self.assertEqual(r['status'],'ambiguous');self.assertIn('Source 1',r['answer']);self.assertIn('Source 2',r['answer'])
    def test_neutral_verifier_uses_exact_evidence_without_unverified_paraphrase(self):
        class Neutral(Models):
            def call(self,op,payload):
                if op=='verify':return {'scores':[{'entailment':.05,'neutral':.94,'contradiction':.01}]}
                return super().call(op,payload)
        m=Neutral();m.draft['claims'][0]['text']='The application charge is twenty-five pesos.'
        r=self.engine(models=m).answer('Application fee?')
        self.assertIn('“Applications cost 25 pesos.”',r['answer']);self.assertNotIn('twenty-five',r['answer'])
    def test_identical_passages_keep_separate_source_provenance(self):
        context=select_context([row(),row(2,doc=2)],Settings('',''))
        self.assertEqual({r['document_id'] for r in context},{1,2})
    def test_larger_candidate_configuration_respects_model_batch_limit(self):
        from dataclasses import replace
        class ManyQueries(Models):
            sizes=[]
            def call(self,op,payload):
                if op=='rerank':
                    self.sizes.append(len(payload['pairs']));self_assert=len(payload['pairs'])<=128
                    if not self_assert:raise PipelineFailure('rerank','oversized_batch')
                result=super().call(op,payload)
                if op=='generate' and 'search_queries' in payload.get('schema',{}).get('properties',{}):
                    plan=json.loads(result['text']);plan['search_queries']=['fee','application charge','application cost','registration payment'];result['text']=json.dumps(plan)
                return result
        models=ManyQueries();rows=[row(i,doc=i,text='Applications cost 25 pesos.' if i==1 else f'Other record {i}.') for i in range(1,49)]
        result=Engine(replace(Settings('',''),rerank_count=48),DB(rows=rows),models).answer('Application fee?')
        self.assertEqual(result['status'],'answered');self.assertEqual(models.sizes,[128,64])

class IngestionTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from transformers import AutoTokenizer
        cls.tokenizer=AutoTokenizer.from_pretrained(ROOT/'models'/'embedding',local_files_only=True)
    def test_long_bilingual_source_preserves_text_and_offsets(self):
        text='\n\n'.join(f'Paragraph {i}: Applications require identification. Ang serbisyo ay bukas tuwing Lunes.' for i in range(100))
        chunks=chunk_document({'title':'Varied source','sections':[{'title':'Access','content':text}]},self.tokenizer)
        self.assertGreater(len(chunks),5)
        for c in chunks:
            p=c['provenance'];self.assertEqual(c['content'],text[p['char_start']:p['char_end']].strip());self.assertLessEqual(c['token_count'],512)
        self.assertTrue(text.strip().endswith(chunks[-1]['content'][-30:]))
    def test_idempotent_chunking(self):
        d={'title':'Policy','sections':[{'title':'One','content':'A complete source paragraph.'}]}
        self.assertEqual(chunk_document(d,self.tokenizer),chunk_document(d,self.tokenizer))
    def test_summary_is_not_source(self):
        with self.assertRaisesRegex(ValueError,'source_text_unavailable'):
            source_units({'summary':'An unsupported summary.'})
    def test_path_traversal_rejected(self):
        with self.assertRaises(ValueError):source_units({'file_path':'../../outside.txt'},ROOT)
    def test_empty_source_fails(self):
        with self.assertRaises(ValueError):chunk_document({'sections':[{'content':''}]},self.tokenizer)
    def test_docx_tables_preserve_order_and_values(self):
        from docx import Document
        with tempfile.TemporaryDirectory(dir=ROOT/'tests') as folder:
            d=Document();d.add_paragraph('Before table.');table=d.add_table(rows=1,cols=2)
            table.cell(0,0).text='Service';table.cell(0,1).text='25 pesos';d.add_paragraph('After table.')
            d.save(Path(folder)/'source.docx')
            text=source_units({'file_path':'source.docx'},folder)[0]['text']
            self.assertLess(text.index('Before'),text.index('25 pesos'));self.assertLess(text.index('25 pesos'),text.index('After'))
            self.assertEqual(text.count('25 pesos'),1)
    def test_pdf_without_text_is_explicit_failure(self):
        from pypdf import PdfWriter
        with tempfile.TemporaryDirectory(dir=ROOT/'tests') as folder:
            writer=PdfWriter();writer.add_blank_page(width=300,height=300);writer.write(Path(folder)/'blank.pdf')
            with self.assertRaisesRegex(ValueError,'no_extractable_text'):
                chunk_document({'file_path':'blank.pdf'},self.tokenizer,folder)

class WorkerContractTests(unittest.TestCase):
    def test_model_mismatch_is_not_published(self):
        from worker import process
        calls=[]
        class JobDB:
            def rpc(self,name,payload):calls.append((name,payload))
        process(JobDB(),None,None,{'job':{'job_id':99,'revision':1,'lease_token':'fixture','model_id':'incompatible-model'},'document':{}})
        self.assertEqual([name for name,_ in calls],['bb_fail_job'])
        self.assertEqual(calls[0][1]['p_code'],'model_version_mismatch')

if __name__=='__main__':unittest.main()
